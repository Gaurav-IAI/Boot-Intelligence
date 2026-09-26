"""ECI statistical report — "Detailed Results": candidate-wise votes per assembly constituency.

Two published forms are read:

* Excel (2023 onwards): one row per candidate with STATE, AC NO., AC NAME, CANDIDATE NAME,
  SEX, AGE, CATEGORY, PARTY, SYMBOL, GENERAL, POSTAL, TOTAL, % VOTES POLLED, TOTAL ELECTORS.
* PDF with a text layer (2018): per constituency a "Constituency <n>. <name> ... TOTAL
  ELECTORS : <n>" line, one line per candidate (name and symbol may wrap onto following
  lines), and a closing "TURNOUT ... TOTAL: <general> <postal> <total> <turnout%>" line.
  Columns are taken from word positions under the header, not from text order.

Every number is read from typed text or cells; nothing is OCR'd.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

PARSER_NAME = "eci_detailed_results"
PARSER_VERSION = "1.0.0"


@dataclass
class Candidate:
    serial: int | None
    name: str
    sex: str | None
    age: int | None
    category: str | None
    party: str | None
    symbol: str | None
    general: int | None
    postal: int | None
    total: int
    pct: float | None


@dataclass
class AcBlock:
    ac_number: int
    ac_name: str
    total_electors: int | None = None
    candidates: list[Candidate] = field(default_factory=list)
    printed_total: int | None = None          # PDF "TOTAL:" line

    @property
    def total_votes(self) -> int:
        return sum(c.total for c in self.candidates)


def _int(v) -> int | None:
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return int(v)
    s = str(v).strip().replace(",", "")
    return int(s) if re.fullmatch(r"-?\d+", s) else None


def _float(v) -> float | None:
    try:
        return float(str(v).strip())
    except (TypeError, ValueError):
        return None


# ------------------------------------------------------------------------ Excel
def parse_detailed_xlsx(path: Path) -> list[AcBlock]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    head_i = next(i for i, r in enumerate(rows)
                  if r and any(str(v).strip().upper() == "AC NO." for v in r if v is not None))
    head = [str(v).strip().upper() if v is not None else "" for v in rows[head_i]]

    def col(name: str) -> int:
        return head.index(name)

    c_ac, c_acn, c_name = col("AC NO."), col("AC NAME"), col("CANDIDATE NAME")
    c_sex, c_age, c_cat, c_party, c_sym = col("SEX"), col("AGE"), col("CATEGORY"), col("PARTY"), col("SYMBOL")
    c_gen, c_post, c_tot = col("GENERAL"), col("POSTAL"), col("TOTAL")
    c_pct = next(i for i, h in enumerate(head) if h.startswith("% VOTES"))
    c_el = col("TOTAL ELECTORS")
    blocks: dict[int, AcBlock] = {}
    for r in rows[head_i + 1:]:
        if not r or _int(r[c_ac]) is None or _int(r[c_tot]) is None:
            continue
        n = _int(r[c_ac])
        b = blocks.setdefault(n, AcBlock(n, str(r[c_acn]).strip(), _int(r[c_el])))
        raw = str(r[c_name]).strip()
        m = re.match(r"^(\d+)\s+(.*)$", raw)                       # "1 Dr.palvai Harish"
        serial, name = (int(m.group(1)), m.group(2)) if m else (None, raw)
        party = (str(r[c_party]).strip() or None) if r[c_party] is not None else None
        if party == "NOTA" or name.lower() == "nota":
            name = "None of the Above"
        b.candidates.append(Candidate(
            serial=serial, name=name, sex=(str(r[c_sex]).strip() or None) if r[c_sex] else None,
            age=_int(r[c_age]), category=(str(r[c_cat]).strip() or None) if r[c_cat] else None,
            party=party, symbol=(str(r[c_sym]).strip() or None) if r[c_sym] else None,
            general=_int(r[c_gen]), postal=_int(r[c_post]), total=_int(r[c_tot]), pct=_float(r[c_pct])))
    return [blocks[k] for k in sorted(blocks)]


# ------------------------------------------------------------------------ PDF
_COLS = [("name", 50), ("sex", 178), ("age", 203), ("category", 228), ("party", 280), ("symbol", 322),
         ("general", 380), ("postal", 438), ("total", 482), ("pct", 530)]
_AC = re.compile(r"^\s*(\d+)\.\s*$")


def _col_bounds(header_words) -> list[tuple[str, float]]:
    """Left edge of each column from the header words (fallback: the 2018 layout's edges)."""
    x = {w[4].upper(): w[0] for w in header_words}
    if "GENERAL" not in x:
        return _COLS
    return [("name", 50), ("sex", x.get("SEX", 180) - 3), ("age", x.get("AGE", 209) - 6),
            ("category", x.get("CATEGORY", 230) - 3), ("party", x.get("PARTY", 288) - 10),
            ("symbol", x.get("SYMBOL", 324) - 6), ("general", x["GENERAL"] - 6),
            ("postal", x.get("POSTAL", 445) - 8), ("total", x.get("TOTAL", 486) - 6),
            ("pct", x.get("POLLED", 534) - 6)]


def parse_detailed_pdf(path: Path) -> list[AcBlock]:
    import pymupdf
    blocks: list[AcBlock] = []
    cur: AcBlock | None = None
    last: Candidate | None = None
    with pymupdf.open(path) as doc:
        for page in doc:
            # group words into lines by closeness (a wrapped name is centred on its row, its
            # first line starting a point or two above the row's numbers)
            lines: dict[int, list] = {}
            line_y: list[float] = []
            for w in sorted(page.get_text("words"), key=lambda w: w[1]):
                if line_y and w[1] - line_y[-1] <= 3.5:
                    lines[len(line_y) - 1].append(w)
                else:
                    line_y.append(w[1])
                    lines[len(line_y) - 1] = [w]
            header = next((sorted(ws, key=lambda w: w[0]) for ws in lines.values()
                           if any(w[4].upper() == "GENERAL" for w in ws)), [])
            bounds = _col_bounds(header)
            header_y = max((w[3] for w in header), default=0)
            last = None                     # a wrapped line never continues across a page

            def field_of(x: float) -> str:
                name = "serial"
                for f, left in bounds:
                    if x >= left:
                        name = f
                return name

            for k in sorted(lines):
                ws = sorted(lines[k], key=lambda w: w[0])
                if ws[0][1] <= header_y:        # the page title and column header
                    continue
                texts = [w[4] for w in ws]
                if texts and texts[0] == "Constituency":
                    num = next((int(m.group(1)) for t in texts if (m := _AC.match(t))), None)
                    ti = texts.index("TOTAL") if "TOTAL" in texts else len(texts)
                    name = " ".join(t for t in texts[1:ti] if not _AC.match(t))
                    electors = _int(texts[-1]) if ":" in texts else None
                    cur = AcBlock(num, name, electors)
                    blocks.append(cur)
                    last = None
                    continue
                if cur is None:
                    continue
                if "TOTAL:" in texts:
                    nums = [_int(t) for t in texts if _int(t) is not None]
                    # the constituency's own line comes first; a later statewide "TOTAL:"
                    # (after the last constituency) must not overwrite it
                    if len(nums) >= 3 and cur.printed_total is None:
                        cur.printed_total = nums[2]          # general, postal, TOTAL, (turnout %)
                    last = None
                    continue
                first = ws[0]
                if first[0] < 52 and _int(first[4]) is not None:
                    cells: dict[str, list[str]] = {}
                    for w in ws[1:]:
                        cells.setdefault(field_of((w[0] + w[2]) / 2 if field_of(w[0]) != "name" else w[0]), []).append(w[4])
                    g = lambda f: " ".join(cells.get(f, [])) or None  # noqa: E731
                    total = _int(g("total"))
                    if total is None:
                        continue
                    name = g("name") or ""
                    last = Candidate(serial=_int(first[4]), name=name, sex=g("sex"), age=_int(g("age")),
                                     category=g("category"), party=g("party"), symbol=g("symbol"),
                                     general=_int(g("general")), postal=_int(g("postal")), total=total,
                                     pct=_float(g("pct")))
                    cur.candidates.append(last)
                elif last is not None:
                    # a wrapped name or symbol continues the previous candidate
                    for w in ws:
                        f = field_of(w[0])
                        if f == "name":
                            last.name = f"{last.name} {w[4]}".strip()
                        elif f == "symbol":
                            last.symbol = f"{last.symbol or ''} {w[4]}".strip()
    return [b for b in blocks if b.ac_number is not None]


def parse_detailed_results(path: Path) -> list[AcBlock]:
    return parse_detailed_xlsx(path) if path.suffix.lower() in (".xlsx", ".xls") else parse_detailed_pdf(path)


def check_block(b: AcBlock) -> tuple[bool, str]:
    """The report's own arithmetic: each candidate's general + postal = total, and (PDF) the
    candidates add up to the printed constituency total."""
    bad = [c.name for c in b.candidates
           if c.general is not None and c.postal is not None and c.general + c.postal != c.total]
    if bad:
        return False, f"general + postal differs from total for {', '.join(bad[:3])}"
    if b.printed_total is not None and b.printed_total != b.total_votes:
        return False, f"candidates sum to {b.total_votes:,}; the report prints {b.printed_total:,}"
    if not b.candidates:
        return False, "no candidates read"
    return True, ("candidate votes reconcile with the printed constituency total"
                  if b.printed_total is not None else "general + postal = total for every candidate")
